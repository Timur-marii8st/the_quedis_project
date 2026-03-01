import io
from typing import BinaryIO
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from ..models import db, Group, Student, Subject, Score
from .distributor import ensure_scores_for_student


def _normalize_name(raw) -> str:
    """Берёт первые 2 слова из имени (как в оригинальном проекте)."""
    return ' '.join(str(raw).split()[:2])


def import_students(group: Group, file_obj: BinaryIO, replace: bool = False) -> int:
    """
    Читает .xlsx файл, первая колонка — имена студентов.
    Если replace=True — удаляет существующих студентов перед импортом.
    Возвращает количество добавленных студентов.
    """
    try:
        df = pd.read_excel(file_obj)
    except Exception as e:
        raise ValueError(f'Ошибка чтения Excel-файла: {e}')

    if df.empty:
        raise ValueError('Excel-файл пуст.')

    name_col = df.columns[0]
    names = (
        df[name_col]
        .dropna()
        .apply(_normalize_name)
        .str.strip()
        .unique()
        .tolist()
    )
    names = [n for n in names if n]

    if replace:
        for student in list(group.students):
            db.session.delete(student)
        db.session.flush()

    existing_names = {s.name for s in Student.query.filter_by(group_id=group.id).all()}

    added = 0
    for name in names:
        if name in existing_names:
            continue
        student = Student(name=name, group_id=group.id)
        db.session.add(student)
        db.session.flush()
        ensure_scores_for_student(student)
        added += 1

    db.session.commit()
    return added


def export_scores(group: Group) -> bytes:
    """
    Генерирует .xlsx с таблицей баллов: Студент | Предмет1 | Предмет2 | ...
    Возвращает байты файла.
    """
    subjects = Subject.query.filter_by(group_id=group.id).order_by(Subject.display_name).all()
    students = Student.query.filter_by(group_id=group.id).order_by(Student.name).all()

    score_map = {}
    if students and subjects:
        student_ids = [s.id for s in students]
        subject_ids = [s.id for s in subjects]
        scores = Score.query.filter(
            Score.student_id.in_(student_ids),
            Score.subject_id.in_(subject_ids),
        ).all()
        score_map = {(sc.student_id, sc.subject_id): sc.value for sc in scores}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Баллы'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1D3557')

    headers = ['Студент'] + [s.display_name for s in subjects]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, student in enumerate(students, start=2):
        ws.cell(row=row_idx, column=1, value=student.name)
        for col_idx, subject in enumerate(subjects, start=2):
            ws.cell(row=row_idx, column=col_idx,
                    value=score_map.get((student.id, subject.id), 0))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
